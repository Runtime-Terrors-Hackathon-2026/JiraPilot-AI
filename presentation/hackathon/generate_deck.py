import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette Tokens
    BG_DARK = RGBColor(7, 10, 20)         # #070A14 deep midnight
    SURFACE_CARD = RGBColor(14, 22, 42)    # #0E162A dark blue surface
    SURFACE_ALT = RGBColor(19, 29, 56)     # #131D38 highlight card
    BORDER_INDIGO = RGBColor(55, 65, 95)   # #37415F subtle card border
    TEXT_WHITE = RGBColor(248, 250, 252)   # #F8FAFC crisp white
    TEXT_MUTED = RGBColor(148, 163, 184)   # #94A3B8 slate grey
    TEXT_CYAN = RGBColor(56, 189, 248)     # #38BDF8 cyan accent
    ACCENT_CYAN = RGBColor(56, 189, 248)   # #38BDF8 cyan
    ACCENT_INDIGO = RGBColor(99, 102, 241) # #6366F1 electric indigo
    ACCENT_GREEN = RGBColor(16, 185, 129)  # #10B981 emerald
    ACCENT_RED = RGBColor(239, 68, 68)     # #EF4444 red error
    ACCENT_AMBER = RGBColor(245, 158, 11)  # #F59E0B warning amber

    assets_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "assets", "images"))
    logo_path = os.path.join(assets_dir, "logo-full.png")

    def add_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return bg

    def add_slide_header(slide, title_text, category="JIRAPILOT AI — HACKATHON 2026"):
        # Header category pill
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.5), Inches(0.35))
        tf = cat_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = category.upper()
        p.font.name = "Segoe UI"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = TEXT_CYAN

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.5), Inches(0.8))
        tf2 = title_box.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        p2.text = title_text
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(24)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_WHITE

    def add_card(slide, left, top, width, height, title, body_lines, accent_color=BORDER_INDIGO, pill_text=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = SURFACE_CARD
        card.line.color.rgb = accent_color
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = Inches(0.25)
        tf.margin_top = Inches(0.22)
        tf.margin_right = Inches(0.25)
        tf.margin_bottom = Inches(0.2)

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Segoe UI"
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE

        if pill_text:
            p_sub = tf.add_paragraph()
            p_sub.text = f"[{pill_text}]"
            p_sub.font.name = "Segoe UI"
            p_sub.font.size = Pt(9)
            p_sub.font.bold = True
            p_sub.font.color.rgb = accent_color

        for line in body_lines:
            p_line = tf.add_paragraph()
            p_line.text = f"•  {line}"
            p_line.font.name = "Segoe UI"
            p_line.font.size = Pt(11)
            p_line.font.color.rgb = TEXT_MUTED
            p_line.space_before = Pt(6)

        return card

    # ==========================================
    # SLIDE 1: Title Slide (Introduction)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1)

    # Large Center Banner Card
    hero = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.2), Inches(10.33), Inches(5.1))
    hero.fill.solid()
    hero.fill.fore_color.rgb = SURFACE_CARD
    hero.line.color.rgb = ACCENT_INDIGO
    hero.line.width = Pt(2)

    htf = hero.text_frame
    htf.word_wrap = True
    htf.vertical_anchor = MSO_ANCHOR.MIDDLE

    p = htf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "AUTONOMOUS INCIDENT TRIAGE & ROOT CAUSE PLATFORM"
    p.font.name = "Segoe UI"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = TEXT_CYAN

    p = htf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.text = "JiraPilot AI"
    p.font.name = "Segoe UI"
    p.font.size = Pt(46)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(10)

    p = htf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.text = "Turn Jira Issues into Actionable Investigations"
    p.font.name = "Segoe UI"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = TEXT_CYAN
    p.space_before = Pt(6)

    p = htf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.text = "Your Autonomous Investigation Copilot"
    p.font.name = "Segoe UI"
    p.font.size = Pt(14)
    p.font.italic = True
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(4)

    p = htf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.text = "JiraPilot AI helps software teams organize incoming issues, inspect log evidence,\ndiscover relevant repositories, and prepare clearer investigation updates automatically."
    p.font.name = "Segoe UI"
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(18)

    p = htf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.text = "Built by Runtime Terrors  •  Hackathon 2026  •  Live Integration with Jira Cloud & EmailJS"
    p.font.name = "Segoe UI"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_INDIGO
    p.space_before = Pt(20)

    # ==========================================
    # SLIDE 2: The Problem
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2)
    add_slide_header(s2, "When an issue arrives, where do developers start?", "THE INVESTIGATION BOTTLENECK")

    add_card(s2, Inches(0.8), Inches(1.8), Inches(2.7), Inches(4.8),
             "1. Missing Context",
             ["Tickets arrive without critical reproduction details.",
              "Missing API endpoints, HTTP verbs, and environments.",
              "Developers must repeatedly ping reporters for missing logs.",
              "Back-and-forth comments cause hours of initial delay."],
             ACCENT_AMBER, "UNPREPARED TICKETS")

    add_card(s2, Inches(3.75), Inches(1.8), Inches(2.7), Inches(4.8),
             "2. Scattered Logs",
             ["Diagnostic evidence is fragmented across multiple server logs.",
              "Files contain tens of thousands of lines of noise.",
              "No standardized correlation between stack trace and ticket.",
              "Manual keyword grep searches waste engineering hours."],
             ACCENT_RED, "UNORGANIZED EVIDENCE")

    add_card(s2, Inches(6.7), Inches(1.8), Inches(2.7), Inches(4.8),
             "3. Complex Topologies",
             ["Single client errors mask deep multi-microservice cascades.",
              "Developers manually trace HTTP 500/503 errors across hops.",
              "Difficult to distinguish root failure from secondary symptoms.",
              "Lack of clear, visual request journey maps."],
             ACCENT_INDIGO, "OPAQUE ARCHITECTURE")

    add_card(s2, Inches(9.65), Inches(1.8), Inches(2.7), Inches(4.8),
             "4. Communication Drag",
             ["Locating the right Git codebase requires tribal knowledge.",
              "Writing diagnostic summaries for Jira is inconsistent.",
              "On-call alerts lack immediate root cause hypotheses.",
              "Junior engineers struggle to find where to begin debugging."],
             BORDER_INDIGO, "MANUAL HANDOFFS")

    # ==========================================
    # SLIDE 3: Why It Matters
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3)
    add_slide_header(s3, "The challenge is not just fixing the bug", "INDUSTRY REALITY")

    # Journey banner
    jb = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.73), Inches(1.4))
    jb.fill.solid()
    jb.fill.fore_color.rgb = SURFACE_CARD
    jb.line.color.rgb = ACCENT_INDIGO
    jtf = jb.text_frame
    jtf.word_wrap = True
    jtf.vertical_anchor = MSO_ANCHOR.MIDDLE
    jp = jtf.paragraphs[0]
    jp.alignment = PP_ALIGN.CENTER
    jp.text = "THE REPEATED MANUAL INVESTIGATION JOURNEY TODAY"
    jp.font.size = Pt(11)
    jp.font.bold = True
    jp.font.color.rgb = TEXT_CYAN

    jp2 = jtf.add_paragraph()
    jp2.alignment = PP_ALIGN.CENTER
    jp2.text = "Read Ticket  →  Ask for Details  →  Find Logs  →  Trace Error Across Services  →  Find Code  →  Share Findings"
    jp2.font.size = Pt(14)
    jp2.font.bold = True
    jp2.font.color.rgb = TEXT_WHITE
    jp2.space_before = Pt(8)

    # 3 Pillars underneath
    add_card(s3, Inches(0.8), Inches(3.5), Inches(3.7), Inches(3.3),
             "Frequent Context Switching",
             ["Engineers are pulled away from active feature sprints to triage unvalidated tickets.",
              "Mental reset time after interrupted debugging impacts productivity.",
              "Repetitive triage creates operational fatigue."],
             BORDER_INDIGO)

    add_card(s3, Inches(4.8), Inches(3.5), Inches(3.7), Inches(3.3),
             "Inconsistent Diagnostic Handoffs",
             ["Different developers write different ticket comments.",
              "Critical evidence like correlation IDs or log line references are lost.",
              "Handoffs between time zones and shifts become high-friction."],
             BORDER_INDIGO)

    add_card(s3, Inches(8.8), Inches(3.5), Inches(3.7), Inches(3.3),
             "Organize, Don't Replace",
             ["JiraPilot AI does not attempt to write speculative production code.",
              "It organizes evidence so engineers start debugging with full facts.",
              "Human engineering judgment remains in complete control."],
             ACCENT_GREEN, "HUMAN-IN-THE-LOOP")

    # ==========================================
    # SLIDE 4: Our Solution
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4)
    add_slide_header(s4, "Meet JiraPilot AI: Your Autonomous Investigation Copilot", "THE PLATFORM SOLUTION")

    add_card(s4, Inches(0.8), Inches(1.8), Inches(3.7), Inches(2.4),
             "1. Ticket Readiness Engine",
             ["Deterministic 100-point scoring matrix.",
              "Detects missing endpoints, environments, logs, and payloads.",
              "Prepares polite clarification comments automatically."],
             ACCENT_CYAN, "AUTOMATED TRIAGE")

    add_card(s4, Inches(4.8), Inches(1.8), Inches(3.7), Inches(2.4),
             "2. Log Telemetry Parser",
             ["Deterministic regex extraction of HTTP status, methods, endpoints.",
              "Correlation ID and line-referenced evidence gathering.",
              "Automatic credential and secret redaction before processing."],
             ACCENT_INDIGO, "EVIDENCE EXTRACTION")

    add_card(s4, Inches(8.8), Inches(1.8), Inches(3.7), Inches(2.4),
             "3. Holographic API Flow Map",
             ["Interactive visual request path from Client to Response.",
              "Instant visual failure node isolation (e.g. HTTP 503).",
              "Synchronized request timeline and line-referenced log viewer."],
             ACCENT_INDIGO, "TOPOLOGY MAP")

    add_card(s4, Inches(0.8), Inches(4.5), Inches(3.7), Inches(2.4),
             "4. Repository Discovery",
             ["Matches affected services against configured Git codebases.",
              "Provides approved, copyable git clone commands.",
              "Eliminates tribal knowledge search for repositories."],
             ACCENT_GREEN, "CODE NAVIGATION")

    add_card(s4, Inches(4.8), Inches(4.5), Inches(3.7), Inches(2.4),
             "5. 7-Section Investigation Report",
             ["Executive summary, failure points, and probable root causes.",
              "Confidence ratings, evidence base, and developer action steps.",
              "Reproduction cURL commands and test case blueprints."],
             ACCENT_GREEN, "STRUCTURED RCA")

    add_card(s4, Inches(8.8), Inches(4.5), Inches(3.7), Inches(2.4),
             "6. Live Multi-Channel Handoff",
             ["One-click automated structured comment posting to Jira Cloud.",
              "Animated responsive HTML triage email dispatch via EmailJS.",
              "Live webhook integration and centralized audit logging."],
             ACCENT_CYAN, "JIRA & EMAILJS")

    # ==========================================
    # SLIDE 5: How It Works
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5)
    add_slide_header(s5, "From a Jira Issue to a Clearer Investigation", "END-TO-END WORKFLOW")

    # Path A: Incomplete Ticket
    add_card(s5, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
             "PATH A: Incomplete Ticket (< 70 Readiness)",
             ["Ticket ingested via Jira Cloud webhook or sync.",
              "Readiness Engine scores ticket across 6 criteria (e.g., Score: 10/100).",
              "Status set to WAITING_FOR_INFORMATION.",
              "System prepares polite clarification comment listing exactly what is needed (Endpoint, Logs, Payload).",
              "Audit trail logged; developer is spared from manual context gathering.",
              "When reporter provides logs or payload, system auto-resumes."],
             ACCENT_AMBER, "CLARIFICATION WORKFLOW")

    # Path B: Ready Ticket
    add_card(s5, Inches(6.9), Inches(1.8), Inches(5.6), Inches(5.0),
             "PATH B: Ready Ticket (≥ 70 Readiness)",
             ["Ingested ticket verified READY (Score: 100/100).",
              "Log Analysis Engine parses raw telemetry & decompresses attachments.",
              "Extracts API endpoints, correlation IDs, timestamps, and stack traces.",
              "Constructs Interactive API Flow Map isolating the exact failure node.",
              "Matches relevant Git repository & approved clone command.",
              "Synthesizes 7-section structured RCA report.",
              "Posts structured investigation comment to Jira Cloud & dispatches EmailJS alert."],
             ACCENT_GREEN, "FULL ROOT CAUSE WORKFLOW")

    # ==========================================
    # SLIDE 6: Feature 1 - Ticket Readiness
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6)
    add_slide_header(s6, "Feature 1: Start With the Right Information", "DIAGNOSTIC READINESS EVALUATION")

    add_card(s6, Inches(0.8), Inches(1.8), Inches(4.0), Inches(5.0),
             "Example: Ticket NOVA-1001",
             ["Ticket Title: 'User provisioning fails intermittently'",
              "Description: 'Users cannot complete signup flow. Please fix.'",
              "Missing Fields Detected:",
              "  ✗  API Endpoint & HTTP verb",
              "  ✗  Application Server Logs",
              "  ✗  Sample Request Payload (JSON)",
              "  ✗  Environment (Prod/Stage)",
              "Initial Readiness Score: 10 / 100",
              "Status: WAITING_FOR_INFORMATION"],
             ACCENT_RED, "UNPREPARED INPUT")

    add_card(s6, Inches(5.1), Inches(1.8), Inches(3.6), Inches(5.0),
             "Readiness Checklist Matrix",
             ["Description: Verified (Summary present)",
              "API Details: Checked against OpenAPI schemas",
              "Server Logs: Validated format (.log, .txt, .gz)",
              "Request Payload: Checked for valid JSON",
              "Environment: Checked against known tiers",
              "Error Pattern: Checked for exception keywords",
              "",
              "Scoring weights are customizable via JiraPilot Policy Service."],
             ACCENT_INDIGO, "EVALUATION ENGINE")

    add_card(s6, Inches(9.0), Inches(1.8), Inches(3.5), Inches(5.0),
             "Automated Clarification Comment",
             ["Generated for Jira Cloud:",
              "'Hello! Our triage engine found a readiness score of 10/100. To initiate deep root cause investigation, please provide:",
              "• API Endpoint & Method",
              "• Server Logs",
              "• Sample Payload (JSON)'",
              "",
              "Once updated, JiraPilot automatically resumes triage."],
             ACCENT_GREEN, "AUTONOMOUS TRIAGE")

    # ==========================================
    # SLIDE 7: Feature 2 - Log Analysis
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_bg(s7)
    add_slide_header(s7, "Feature 2: Turn Raw Logs Into Useful Evidence", "LOG TELEMETRY CORRELATION")

    # Analogy banner
    ab = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.73), Inches(1.1))
    ab.fill.solid()
    ab.fill.fore_color.rgb = SURFACE_CARD
    ab.line.color.rgb = TEXT_CYAN
    atf = ab.text_frame
    atf.word_wrap = True
    atf.vertical_anchor = MSO_ANCHOR.MIDDLE
    ap = atf.paragraphs[0]
    ap.alignment = PP_ALIGN.CENTER
    ap.text = "THE FLIGHT RECORDER ANALOGY"
    ap.font.size = Pt(11)
    ap.font.bold = True
    ap.font.color.rgb = TEXT_CYAN
    ap2 = atf.add_paragraph()
    ap2.alignment = PP_ALIGN.CENTER
    ap2.text = "\"Logs are like the flight recorder of software. JiraPilot AI extracts the critical telemetry without the noise.\""
    ap2.font.size = Pt(13)
    ap2.font.italic = True
    ap2.font.color.rgb = TEXT_WHITE
    ap2.space_before = Pt(4)

    add_card(s7, Inches(0.8), Inches(3.2), Inches(3.7), Inches(3.6),
             "Raw Telemetry Input",
             ["10,000+ lines of application logs.",
              "Contains debug lines, heartbeats, and stack traces.",
              "Decompresses .gz and .log.gz files automatically.",
              "Enforces safety policy limit (up to 50,000 lines)."],
             BORDER_INDIGO, "INPUT LOG")

    add_card(s7, Inches(4.8), Inches(3.2), Inches(3.7), Inches(3.6),
             "Extracted Telemetry Node",
             ["Endpoint: POST /api/v1/resources/allocate",
              "Status: HTTP 503 Service Unavailable",
              "Correlation ID: CORR-DEMO-1001",
              "Timestamp: 2026-10-10 10:24:32",
              "Service: resource-inventory-service",
              "Stack Frame: JpaResourceRepository.save():94"],
             ACCENT_CYAN, "ISOLATED FACTS")

    add_card(s7, Inches(8.8), Inches(3.2), Inches(3.7), Inches(3.6),
             "Evidence vs Hypothesis",
             ["Observed Fact: HTTP 503 and SocketTimeoutException recorded on line 82.",
              "Possible Cause Hypothesis: Downstream socket read timeout after 5000ms boundary breach.",
              "Confidence Score: 88% (HIGH).",
              "Always labeled as a hypothesis for developer review."],
             ACCENT_AMBER, "EXACT EVIDENCE")

    # ==========================================
    # SLIDE 8: Feature 3 - Interactive API Flow
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_bg(s8)
    add_slide_header(s8, "Feature 3: See Where the Request Journey Breaks Down", "HOLOGRAPHIC TOPOLOGY VISUALIZATION")

    add_card(s8, Inches(0.8), Inches(1.8), Inches(3.7), Inches(5.0),
             "Visual Request Topology",
             ["Interactive holographic canvas inspired by futuristic observability consoles.",
              "Client Node: Initial ingress point with timestamp.",
              "Trunk Path: API 1 (/auth/login) → API 2 (/user/profile).",
              "Top Branch: Healthy service mesh (API 3 → API 5 → API 7).",
              "Bottom Branch: Isolated failure point (API 4 → API 6 fallback).",
              "Response Node: Final egress status."],
             ACCENT_CYAN, "INTERACTIVE MAP")

    add_card(s8, Inches(4.8), Inches(1.8), Inches(3.7), Inches(5.0),
             "Color-Coded Status Language",
             ["Cyan / Blue Nodes: Verified successful APIs (HTTP 200) supported by telemetry.",
              "Red Glowing Hexagonal Node: Confirmed error evidence (HTTP 503) with pulsing aura.",
              "Amber Indicators: Inferred dependencies or degraded timeout warnings.",
              "Dashed Lines: Explicitly distinguished inferred architectural relationships."],
             ACCENT_INDIGO, "INTUITIVE STATUS")

    add_card(s8, Inches(8.8), Inches(1.8), Inches(3.7), Inches(5.0),
             "Synchronized Investigation",
             ["Click any node to open the Right Details Panel (Overview, Line Logs, Payload, Related).",
              "Synchronized Request Timeline: 8-step sequence synchronized with node selection.",
              "Zoom, Pan, Fit-to-View, and Holo FX toggles.",
              "1-Click Accessible List View for keyboard & screen-reader navigation."],
             ACCENT_GREEN, "DEEP DIVE CONSOLE")

    # ==========================================
    # SLIDE 9: Feature 4 - Repository Discovery & Reports
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_bg(s9)
    add_slide_header(s9, "Feature 4: Move From Evidence to the Next Debugging Step", "CODE DISCOVERY & REPORTING")

    add_card(s9, Inches(0.8), Inches(1.8), Inches(3.7), Inches(5.0),
             "Repository Matcher",
             ["Matches extracted services against configured Git codebases.",
              "Example: resource-inventory-service → https://github.com/northstar/resource-inventory-service.git",
              "Generates approved, copyable git clone commands.",
              "Eliminates guesswork about where the code lives.",
              "Clearly notes: Matching a repository is a starting point, not proof of root cause."],
             ACCENT_GREEN, "GIT CODEBASE")

    add_card(s9, Inches(4.8), Inches(1.8), Inches(3.7), Inches(5.0),
             "7-Section Jira Comment",
             ["Posts structured comment to Jira Cloud:",
              "1. Executive Summary",
              "2. Failure Point & Root Cause",
              "3. Diagnostic Evidence Base",
              "4. Recommended Remediation & Fix",
              "5. Developer Action Items",
              "6. Reproduction cURL & Postman",
              "7. Approved Clone Command"],
             ACCENT_INDIGO, "JIRA CLOUD")

    add_card(s9, Inches(8.8), Inches(1.8), Inches(3.7), Inches(5.0),
             "EmailJS Live Triage Alerts",
             ["Dispatches responsive, dark-mode animated HTML triage emails.",
              "Live pulsing radar beacon and telemetry breakdown matrix.",
              "macOS-style terminal reproduction box.",
              "Deep-link directly to JiraPilot incident workspace.",
              "Verified live with EmailJS API (Service: service_1ooa8fh)."],
             ACCENT_CYAN, "EMAILJS DISPATCH")

    # ==========================================
    # SLIDE 10: Technical Architecture
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    add_bg(s10)
    add_slide_header(s10, "How We Built It: Scalable Micro-Architecture", "SYSTEM ARCHITECTURE")

    add_card(s10, Inches(0.8), Inches(1.8), Inches(2.7), Inches(5.0),
             "React Frontend",
             ["Framework: React 18 + TypeScript + Vite",
              "Styling: Tailwind CSS design tokens",
              "Visualization: Recharts + SVG Holographic Flow Canvas",
              "State: React state + zero external bloat",
              "Theme: Dark/Light mode support",
              "Port: 3000 (Hot Reload)"],
             ACCENT_CYAN, "PORT 3000")

    add_card(s10, Inches(3.75), Inches(1.8), Inches(2.7), Inches(5.0),
             "Spring Boot Backend",
             ["Framework: Java 21 + Spring Boot 3.2.5",
              "Database: H2 (Dev) / PostgreSQL (Prod) via JPA",
              "Responsibilities: Workflow coordination, Jira Cloud client, EmailJS client, audit logging",
              "Port: 8080 (Tomcat REST API)"],
             ACCENT_INDIGO, "PORT 8080")

    add_card(s10, Inches(6.7), Inches(1.8), Inches(2.7), Inches(5.0),
             "FastAPI AI Service",
             ["Framework: Python 3.12 + FastAPI + Uvicorn",
              "Log Engine: Regex telemetry parser & evidence correlation",
              "Repository Matcher: Service-to-Git mapping",
              "AI Provider: Optional NVIDIA NIM LLM adapter with fallback",
              "Port: 8000"],
             ACCENT_GREEN, "PORT 8000")

    add_card(s10, Inches(9.65), Inches(1.8), Inches(2.7), Inches(5.0),
             "External Integrations",
             ["Atlassian Jira Cloud: REST API v3 (Comments, Webhooks, Sync)",
              "EmailJS API: v1.0 REST API with browser-compatible headers",
              "NVIDIA NIM: Optional Kimi-k3 / Nemotron LLM endpoint",
              "Zero-Trust: Redacts secrets before any external dispatch."],
             BORDER_INDIGO, "INTEGRATIONS")

    # ==========================================
    # SLIDE 11: Business Value & Limitations
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    add_bg(s11)
    add_slide_header(s11, "Why JiraPilot AI Matters: Value, Differentiation, & Honesty", "BUSINESS IMPACT")

    add_card(s11, Inches(0.8), Inches(1.8), Inches(3.7), Inches(5.0),
             "Impact by Stakeholder",
             ["For Developers:",
              "• Less time spent checking tickets and grepping logs.",
              "• A clear starting point and copyable reproduction commands.",
              "",
              "For Teams:",
              "• Consistent 7-section investigation summaries.",
              "• Faster onboarding and smoother shift handoffs.",
              "",
              "For Organizations:",
              "• Standardized investigation workflows.",
              "• Reduced MTTR (Mean Time to Resolution)."],
             ACCENT_CYAN, "MEASURABLE VALUE")

    add_card(s11, Inches(4.8), Inches(1.8), Inches(3.7), Inches(5.0),
             "What Makes It Different",
             ["Unified Pipeline: Connects ticket readiness, log parsing, topology visualization, repo discovery, and handoff in one tool.",
              "Deterministic First: Core extraction relies on verified logs, not hallucinated AI guesses.",
              "Human-in-the-Loop: Empowers engineers; does not attempt risky autonomous code edits.",
              "Non-Invasive: Integrates with existing Jira Cloud and Git repositories."],
             ACCENT_GREEN, "KEY DIFFERENTIATOR")

    add_card(s11, Inches(8.8), Inches(1.8), Inches(3.7), Inches(5.0),
             "Honest Limitations",
             ["Incomplete Logs: Partial logs produce partial investigations.",
              "AI Hypotheses: Suggested root causes must be verified by engineers.",
              "Inferred Edges: If trace headers are missing, relationships are marked inferred.",
              "No Auto-Fixing: Does not execute code or modify production configurations.",
              "Requires Setup: Jira tokens and repo catalogs must be pre-configured."],
             ACCENT_AMBER, "HONEST LIMITATIONS")

    # ==========================================
    # SLIDE 12: Closing Slide
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    add_bg(s12)

    # Hero card
    c_close = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.2), Inches(10.33), Inches(5.1))
    c_close.fill.solid()
    c_close.fill.fore_color.rgb = SURFACE_CARD
    c_close.line.color.rgb = ACCENT_INDIGO
    c_close.line.width = Pt(2)

    ctf = c_close.text_frame
    ctf.word_wrap = True
    ctf.vertical_anchor = MSO_ANCHOR.MIDDLE

    p = ctf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "THE BOTTOM LINE"
    p.font.name = "Segoe UI"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = TEXT_CYAN

    p = ctf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.text = "Less Repetitive Triage.\nMore Time to Solve."
    p.font.name = "Segoe UI"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(8)

    p = ctf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.text = "JiraPilot AI — Your Investigation Copilot"
    p.font.name = "Segoe UI"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_INDIGO
    p.space_before = Pt(6)

    # 3 Takeaway bullets
    p = ctf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.text = "1. Organize the Issue   •   2. Understand the Evidence   •   3. Decide the Next Step with Confidence"
    p.font.name = "Segoe UI"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = TEXT_CYAN
    p.space_before = Pt(16)

    p = ctf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.text = "\"Our goal is not to replace software engineers.\nIt is to help them spend less time organizing an investigation and more time solving the problem.\""
    p.font.name = "Segoe UI"
    p.font.size = Pt(14)
    p.font.italic = True
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(16)

    p = ctf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.text = "Thank You!  •  Questions & Live Demonstration"
    p.font.name = "Segoe UI"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p.space_before = Pt(20)

    # Save PPTX
    out_path = os.path.join(os.path.dirname(__file__), "JiraPilot-AI-Hackathon-Deck.pptx")
    prs.save(out_path)
    print(f"PowerPoint Deck saved successfully to: {out_path}")

if __name__ == "__main__":
    create_deck()
